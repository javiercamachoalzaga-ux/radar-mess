# -*- coding: utf-8 -*-
"""
SAIV | Radar Comercial MESS  (v2: pipeline vivo + Laboratorio IA)

Secrets (.streamlit/secrets.toml):
    mi_contrasena  = "tu_contraseña"
    gemini_api_key = "tu_api_key"
    gemini_model   = "gemini-3.8-flash"   # opcional; cambia aquí el modelo sin tocar el código
"""

import base64
import csv
import hmac
import io
import re
import unicodedata
from datetime import datetime, date, time, timedelta
from urllib.parse import quote

import altair as alt
import numpy as np
import pandas as pd
import streamlit as st

try:
    import google.generativeai as genai
except Exception:
    genai = None

try:
    from docx import Document
except Exception:
    Document = None

# ---------------------------------------------------------------------------
# LOGOTIPOS GRUPO MESS (incrustados; azul para fondos claros, blanco para fondos azules)
# ---------------------------------------------------------------------------
_LOGO_BLUE_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAWgAAABqCAMAAACmhDtYAAAAYFBMVEUATGsATpgARIEARIEADXsAQ34ARH4ARH4ANJcA//8A"
    "Nz8Af/8A/wAAP3sAfwAAUJUARYAAAAAARH4APHkAAP8Af38ARYEARXoAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAD+vLKCAAAA"
    "GHRSTlMZEqZiB1rXlAcBBAIBRAL//gD6CwEC0RVA4xxPAAAS0UlEQVR42u1dC3eduo6WX2Gn584Mxq///0/HkmwwYDYkTcrO"
    "ufXqSltHGPNZyHpZwPjHmrTerpqS4zhJGMZpHMG6t/x31CqOoxP5F5leByux09mEAwwiXzBGpWkgbR0OkDuktmBfsOEMY342"
    "KfKE4S6g/6Nl6Y/4owI94H+S6wCd/474Mw+U8vXBijgPbfJ/X7CJGBUAyHGwcBfQIeM8TeU3laNr6wLdtCgyC+eZTzQC/nhB"
    "pL11o3TaOWdFfpabgP7HGgI4yrgFOvPtMdCR10VZjS/ANM5cDi+HdJ67zE+ZJ4go38bRlgQHCgEzzqKDmBUcQOoBPY0DJCDZ"
    "ElO+jDjZWYhMHF4NZxhFUKMCk/lCyruAdhEZMVkn7NDI6FHQb7tA550Sf1VE+oBAG/sOVtOrIV6MocGq6K3KM81AazPeA7Qv"
    "u1ted+dMIzriA/Je3QU6C4wAVliS7XV5HHjhCHq31Wnu3ghRcvDrmtUkmO4CmqRwRidztGm1DlTddAU67yMtR2dZ54B6RmDR"
    "Yyw8PHO7ezWNI2uvCHTe9MF6Md0lOhJL1iwMxNgC7SrQGbsqYebNELXlKtIV4i0R3oFXxb+ejHZZRqfMEpm7401As0xAY0Oy"
    "ftYC7YsglsYQVdRlMxwHI8eyQEAdb4aNlrzfvJrNEh5RZjmdAozKjeY2oAmoaa1HF6BZEI+zkm0q2/P/mZoZmQky8u791fRo"
    "fERlE7++t+nReRqmgDYDHWeg398dq9cTG4KBgZ7qn2ye+MCWJdssWd17OaDzI6pRgnNJReLrm/RoZw2b0HJo9GgGGjdLOV+W"
    "VWUGWio2WqoZWAYYo0qvaIAjT8vx7S3PD9W9G51KVitlSCNGjVoIY4zQVS2xzpj8a5PIO0NAD1bnrkyi62KBwhGcnXteDWl8"
    "RCCF6DbRoWRuuIHprnbmYdH8Z6Bh7ikalF13vB7SYZ7ojVoHKQtKRZTEaEEHAGjsaA0iN9AtR0Pua1ENSAPevm7TAYK3dwKd"
    "YrMVjvqpq2IB2v7Ydp/WAbGiPMYT39tfoH9jMwwYfmClIesV+iRUUU3wv0B/XOvQKIbBACR77uLMKohRxtq/QH9CvfNVaxDe"
    "/vvbrXq0D3lPDpdgBgEi/AX6k0D/N7W/QP/bgRb7tpYqpfOJZPG60Fw1WTwbQXyNvrrs81RoNv7kDuHgDi/J0bAOAHaFs1/3"
    "ngpwD9C5z9mM9H7cPNDBVWE/XJivvy04q/ateO4QZZ2gdEJauTVaV8hMo5I78XiQJW9dMkjNP7RY+1S6AdbVVPCi5A7WlQdK"
    "zrTEesH/xpjhrpF9iD4kJWN7HfmeYIOAMy1RRBqvnwDmzCC3E5LkG4QnV8HqLnwnqQQJlC3MTqg9MfkWbfC3AZ3Qq79qBHSe"
    "fSoG41h62ePsGqZGIlnDLzORVAeYIQZKjs2gzchS6b69pLWF3VXzoyixvYVYTbshjpLc5bemG7SNgBb2F813mtrQFcVQVBNf"
    "1mpDU3LyusmOYDVHCKZpc8cystE7yUR2q1mcMevL+GlMewsnu3eYHx1eC2gKSewfrQRmHYergNm5C8A+Ba8AtoVgtYrR7CIH"
    "DN3BRXzV4nkJc6znsL0Y0CWXYOwjLRMOEJ5O2Wx4OrPzwdKtp7Zx00JJNHvWZk8YlDzCJ3MfbHgloIV5Ml/K4gienmo6plnz"
    "tLY6HvLlZhF161nUJzBjnqWYU5KeEyMLZQ3+hYCWZyMM1pV8jmcvaWikOcQrM5solhaa5TnFOVaGfrryM0OL1xEd4ziev+OK"
    "AjNPH0qGGbB/TgFbI1cvDC6eXDbNMkqf4jxi5mZ4MaBPGQ/khcUIc1wmXr4J5uBAFRzDKXksRqD3VxflJwG9pDY9BawCHX4N"
    "h8pZT3APfGHoS6dppbiYIqGDHbpvZkucF8X/MKAvXEIpIoGFZw/n8anK8g9dt39rtstTs7HD5szH+iZ88KNsmz8L6EvXsPjs"
    "aQ5k1RhhjFBKdsR0yqxH6WfT8V0JvSqhMed42k9x1RXTjU6lZ0Djk5xoZPzr6VB2IAr+ETt4yUXNBtkRpngKSXWuU5CcS1Cc"
    "GVnPrF7Rzk2kAeEcKLPKc30xoKcrvDs9oylCWuyNH9KyrRaADd/mvTqZ6iGDTdNNlFi2Ehr2WqpsiNXbGOH9BTl6Ih+MUYN8"
    "DrNEX+eRnRDpPi52hcqCAuw2MmRpZ23sqQ2BIwD4Fyg36+mmoxjxTTQ5u51SVd9+MdFRHwLkEzOuvP+pT/OW0IcJ+3da2cfK"
    "DeLTFlOUp52ZGQfbdLpirOzETL77EoTx4u4ISx/oCVOWrMjvNmaMyiOcFcXEAUNfQ1/ZtkF0BEB0G8/Rrz1OyvZEjmrfBBtg"
    "e25hdXe9Sp+A8IJAU5a0r2/lkSWgbN2KRHCyt+njeQDXMYS9hlXL21vcGcsdJuXDM2EfwurKczpXA7vA2UsBTSJycaD1bTK3"
    "MIxoj1ysge7A1YsKbMnopGhvFxXO7sHuAj1y3AbBDq8KtGtjzMLKc0fofssjoB+dd3oysOHo3NTOg9hZO5a9UgkC0IXnQNNm"
    "nVU8vU7Ceh2gt8eMRfcljm6dkyo7Si9ytPykPZQF03SkuKO2Q5mCs9ztO9CZOKtP0ISVXwnoNbNC9yVeL4bbW2YEtE3xU1PE"
    "8W186mbB6iB1nsfOz3lhUiV+JdFxBWhlnV3rVz3Rcd0/ulfCrTxzAWT18j/lUdKZAyyq4iCHC6vcNcmmHuH0W0CLFdB9F67a"
    "0Ig+0J/mnze9X7s9Hsrys4TnIqpm2cOnOPrA0TidWs6nQK8Dd6GcF9/70DdHATqxxyfBx1OPt32ce5hRh/NPVKOt4X9t6ePc"
    "FnjjE6/a9Cmgpy3QuptkI1bx6qMg72eBZiVQjBeQZq0inO+6JZB5DrRyqbRyyFJiln6C/5FxOSSMsgg7jTqOHl+Igp8DDRug"
    "xdcD7S5Ep6Z6oOY8wlL8HxeAbtX7sXVOzTlFlEk0e6yOZvklQIvvBjqbpGCvpCjw1q273qst0m/uQrqB8g9P7eEVOSOyeRS8"
    "DmhkFhMYrU4BQXMnHKzxzUBP03QNaOCI62kcHEj1CMW3NZ2o51c42lWtcYjZcNKtGktIZ5xd23kQqf4ZHG3KUV3Z2fC7ClCe"
    "s4orVbu3zz0+AHTmVbfJUgv2IXdpPu5g9W4FekIf9pWWatYiDHFjq+xsm/eadTYnUR6ujPkQ0HZ/bD6LKNgmrh0YwH8SaOix"
    "6gfPy6PXvuTh9gMMzjaH0pdc4y7Yw4eA7p5G2BczCh7u5ujUYUABV1qTrE9PVzKke76jRqfnFx2MOTIr5Qc5Gs8KCOgcQAgr"
    "mdLVef4g0KEpC9n3klw6vFIcnQ/oZYtOc+jQth7UZDoqC26dHwM6VKnUPuz+mEHfT/vngH50JvDWOw+g/brp9jEbsLXq7ptu"
    "TVzA1r2n/yhH5zdpk+2OtrKRUoUzD+cfBbozAeE/UD4FtqeAOuBVZMRqSiC6vsMPAR2Kw0Xq0OKslhBOFdvqXhntunkA2/MA"
    "3j6SW7eUECo8b+E3h76626vTFmTmO7favEQnR2r6CNAwF5xqspA5twHHHebO24GG0MuGhI1y6rMVsOa9SGq04QMxbZCV9vdp"
    "J6MdFXhBFgvQrora0X5ERvvq16L8tgWP+ZFSRfJ2oEMnIkB+5EYKZnHaCaInq9mylUYvcb+wd3yT1lFNBuTqLGDI0eSF63D0"
    "R7SONpwmFuZ4zI8678L3A92LsUzMegJCCJjQYGXPCBE0dz6f1KSB7DKAueAWcRmn9DTzMr+n3rVcMnf+M486LUHsFwD6IOLY"
    "FG7rpuhknueqvyWWYhSmjyRQHX1V6hrK4tNdWUugA42dFKrpQ5ZhmwC1Qn/xnLwOR3vfyULnjDMsgAU9ywKPC+zzQ6sjvuMo"
    "EqGKzTmfN771LUn9IaDV7H/uAv1SoqMfoT4JUOQneMRLDr9nvqupJ5K+AujpBTmaE9GfoDB13USul7Z4FC8PupMR1r/AfAHQ"
    "6kWBtvrIJX/IqJi0dy1PgY60nYVx2104/IuBDqdH5XahPd3dQrttuHB466OhrB8KdPmEyGWcB45ZXKOOWlxk6In1XjjNEvi5"
    "QNsPncAesB6LusijUVt9bVGmiU1mqLObvgvoeew/DzQfnb0WKzR5b+tmkhzhfB5XrG+Knguj4P++i6OXo8c3AF2rFFx5uzEu"
    "9TBxHC+lz+AHEuQV2prVBGPJoj9MBPlt0aGKnnUH0HPo9IlGzIc16tTVU6hLifZ6LuUkgsslKnyT5GhoC/0eoKUVnHlzC9DW"
    "/1PBm7rnYUcuuhK21WqmQ5ihUGN5H5AHAfDSI+dwNjAE0HV3/T7Q+K0V7zi0dQ/Qq1I/TQmepkTQqqI6nXKVcU29/AsLz4T2"
    "4FATAG+Ii2fFLREEIAs/L076Lo7GoiBU9nwL9N4f8Lncuw7NJri2jlK3foxB7QpjEbUDuXeGRHbRrSYJHCaMPeLUjgx4MpcL"
    "tXyLjCb99J3UyDXQ8W3b4g7ouCdKW47e00Sxj2LWiLaRUtIVUcrB0FmJTkG2OQA+ZGomztRc7yvs6u/5ZuhMi8SKP1bQxrgA"
    "GQmaj0R8gx6tPDni14W659TJubntE/do9AUa0a/+Nx+IwEv6ZQpXkb/yLJnaHZcb5LBA7cbbz18gWB/iAg7wiKPEHaOEBk3t"
    "4Rf/kfG6tFXng/tAC2WaQX5RzYB4T1FQP2ejeHL610B3gJNSmxpEqAz8jFrjDWhoJtBIvFsRMPb657waTE86e2Flszq36vet"
    "A9I5DXX6OYsq//u0MKf3H6lG+yHqY2KQ9FGeACmJg1ZfSZHkgunypi5AD2mhbS4HrmcJg/g+tqX4dRYae6aB5NK3fqUFY+cX"
    "lgLGAYtU+Isev081wJoBqRyG/45GsWi0udbJQ7SjZm0qiud3vsazfSr6KMTZDcpmOFjxrp38NpyN/V96XnkD0Kz/pe+rXE3H"
    "laK7BDTabv7okPtX4IzeGjn+EaDVDmhxDjSKutN76L46Q07va0DzsSHHn1L6+hYtgCct/RuAxjoIWgi352isf5J1NMdAQz1A"
    "7rNywBoBf6IIf+atJ9Z64Fm3sJzFWRWHQB8XoE8sRgj5ZoHLW5eRFqCzkldzUVEF8UhZqo9rDAfXY0PpuzhaWsdlK7+z9PwW"
    "6KbI4FjSfRqjbqOwx6b+5VLSlKkL2lyHOS054mGuEsFA1yQ5b49q/hYvoZafPGp6RXbIrffuq5oxKgn8gtwG6LwnqAF/JfDL"
    "16MwaA8HNDazsWfwn0KYzHFJZHsPqz2WM/KQe60wyln8QgwDpIzBVHvkyGxKIgGgcY2JnQZmjkZCiQlLqGtmq0S4PLRQgKIi"
    "TyazdPVIfNdWuPi6vxpoTy8hJl3IDdDvHMtGB4QT5R9RlNKutPQFHvpJfdFV74qk3Bg5ssA3xTrmKzGDLkqKAW5GckPN9Azz"
    "zWXiDAYM+Eo4i2R9BdLTdwAdlpQutQb6veauYSHM2euUkU5N/sYCNGdP/7L1681lvvSRZ8QMPUaGj+LH4iwUu5FkuXCpFI2y"
    "Us0VQMR5zPBLoP56oKmu3RjNQIcrV0BzicBBEaMmopKzCKMjp7HlaMRhMFA/yTVQlMWYscBrsGKvZMgVe5VlohOXNJKgMcg7"
    "aYa66Hy5pPGoQ/7c77DwW51IcRy2MpoeDQkYaOCeEOu7ProF6IBokveVgSaPGFeTIMbljLCEEibRtYrrsUjeaFGdiIV1OYeX"
    "JQoPoljc3Fao+0uApkchWLZAa0Q02+NFdMRfDjDnGV3ueXdElAge+uJuBBS3IlSgJ4dnfyX4YXnf8dWnq+iCX4GAhmx7xwbo"
    "zL9BFsEMCcgv/F52CGn1zwa6cNEWaL/haE7RlYi/ajiaNzIGes/RBeiYlWCdFfL0VvgS7/mLORpWHO34LrwDUjb7nDCnriTQ"
    "vCjQpHCNUcm9jOZtkkUoy2glZ+mpWEYjqykMrkZMfe4BTUOSsAEpFX3kGaqKQeFmHsmJ8iPfcNkJlCoWt2OVxv/gb2XNhRin"
    "vYwWVXsgoEsEzy11acxyQE+RHDgCWrG+939yBbReQp6qMPmsdcxBD9Q6NAdLzG0V0b9EjWZjdood0cGq7NvM0aQC4/5Ztd9Z"
    "EZaPI45mOVNsrUQqNgPNpk0sOBfsi2FNejQhTdJLU6kzZOgf/vU3hWnl2QxDWyzxj2IIu/yrhJZZ/r0Q2UxkIzt305ErdF7g"
    "B7MVeSaUKX5suoIMO8s/3y2YYUAql8fJeBpVijrSSGgy4i+gWIaqJocoUWR0KobPv/Uze6En1OfuYO3VY7T++Q30/i6J12/A"
    "B34wQy8JND8V6Kxiafoam8bgnm5C1N45wVHTEIIXi/cuK2W6UAjgeHAITaYGLQJybekNWU1sqGbaZaQ1YdmJI0kaPzP03w9H"
    "fn2rflBTVFBX8jr+VMMN5r+jUQqaNPyV+Vo/5f8BDO8OR3syVScAAAAASUVORK5CYII="
)
_LOGO_WHITE_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAWgAAABqCAMAAACmhDtYAAAAYFBMVEX///8AAAD///////////////////////////8AAAAA"
    "AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAk2V76AAAA"
    "GHRSTlP9AAYu0m6xTY8AAAAAAAAAAAAAAAAAAABdzm5qAAANO0lEQVR42u2djZajKgyAJQHC+z/xlQSQX8HOdLS96znb6VpU"
    "/IghJAG37c82q0Ad2/6drN+Lzv+ICvT+R4PZ/wCE8qAs71SyA/l/BvhE+6fx/9+0w+LMj9l8DTXXe/+63QQaAFzYr3PQvHMA"
    "WutwIk8aFOp0anoiaV9DbRB3SXJyA/eAdsWPAbRsA9Bp03szKcz3mOeR3uu4OQCpqrtLokGRILO6Br3L7Ri0TlxDcX0c/ziB"
    "tna/y72CbieNt4G2QQkE4hG0RgAcgbYIKGo5NhRFDeKeBtrLMYLZCPd/O/N7QEPACp5l3hkGyRzoaP4pFHNRskWHaHgYaVBm"
    "r5PZa0q7ONB2G2im479IJZLVwUpNQPuvOWjDv4V2YRFBLh7b7GGgnQ3Pnak6lL8HLRJtCtBsUPQl2vkOJhCWP+T3wFMl2iov"
    "0c7f242gRbPue3AbgPaaIu0MhQ6VbqKOFxG3T9TRsOtoLxm7dOvbOkOx7rRzWw8018sSxa4vWB3OuGhJS/toI4OW59l3XqU5"
    "RbsuxF13bHQb6OphykGnZkhGcm3vY13E6/PnjVfCQGpHbu8bGcqzPwAN2ejFQ6xAG1bNLh/BPHBo6LtC57tr0hpuGxn6aghM"
    "5xrQXnkkjO7oHaVtdJCTeILk/XjgGNz6AYtuH8k/BO03NMYEi3hHhUSEkFxOQPuvhljAoxnCBxxUwf/f73imU8nXCuWebnQq"
    "Gbdv3ozANBYp+5LsWzKsU/2LvwqeyTlVzP+52eowZNIIWsYqebcd92SGdVEmK6IeizrU7zbQoGsz4tSx+0hb+cp2n9WRuZP1"
    "xPf2D/SPrA4JP/hBy8xo+Af6J1aH74sRCReMBt8o3gL5B/ol8y51YvDB/D4BdLIaLvXe/0C/APr/tP0D/fVWR2cb/HqqTq6q"
    "n6sHqUsHjQs/UaLru+neXbVzSq3HCFaOWjnR2RXgTtDBXqu24LmTDww7EaDDUXbGMhQG6jNeII4qOQZhyrqqSqpNr127V8gu"
    "cW/MsHXni9+uiPxY15jaXMhmhXxa2FCqhRe5OuRvHcHJwwDsX6yDULp7FPCowNi6sGPfov/5YaB9fU0bX2OPM+SYXXt6Ggi1"
    "tMqoUjQYL3mn4fgog80lzCAsqKVhnibRo/rqA6Mv1L9Az2UC43P2GjFDR6cBVUvF7bjTW9+ft0eB5pDEsLLBI7L/GRZqB+lT"
    "YGUjZhWc5Ral3CiA+SUeBlqdxoptCNye5UjUpKeyVjZi17c42OB4wsw2b5QngYZJhTl16bwxKtIrwLJGHMboO5s5OM+bEuFR"
    "oO2CYEwRYJnUsJq4ksfRVzjraBcuyDPX+0mgtxUx0lPZvASsQ7pMdpgI9OwJk0cVPg30hvPHlDLr5Eoilj1k1C0L9HqjfBjo"
    "NWBRMK8lJbto1Sw8B6SuNsr3gd6SGeguHhgTcy4I9Mp9iJb5QtCRF44GdURIxg7wLVXtEGiz0ijqO0G7mFnYH6KlNCnblb0+"
    "O8Mp7xicMIc6713Eem8SpLYMjfKFoK2Ma6griSpzFbueeRAy3WubMdXc+z8OgcauSX6kY+mkZR4HWjtDxk26MWuIzEiVyoSN"
    "nqyRKtKcWtKsTnVP0jNPPqacym5rYl4YKDbK00AbGD7YzfM/sCo0KOjKmimdR53G4EkGnQbKMwFPco+LkUxV9lmgNcZn+8Sp"
    "Y1KZgX2AXWnNEaiRQJLqCWnZQvmIyI2cvQ3rZ4E+KjkebRx3PbKUsX9+14lS6qaOXd2O/UhVF7SFbnLEo0Cbs8kXHbkclMGR"
    "8mw30+qO7pDapMAXzEBLLKEp/CjQlQK0M/fcQOyxT4Cw3Ux7LA476Qh7AnrjIFkIYcETQddZkGbiZRvdKY4aac1tdWrtIOS5"
    "bGfjFS3xrtuj4F2LA867qq2e892/U1QAr87rczNPlISBF92D1qSI5JNA1057mqjxURlUr08J1mru6nCxz14Yd6WJTE8CjXV0"
    "ZKEx8LdBLzmYjYqOq7mKCln2Twb9ehn1+kRVXPJiWzj1XHWEY6WR05bt+2LQS/BsnMbkFkkvPCaHee9i14uHJyvpItl5kkLx"
    "l6DNTyR66ejoI1zqdnHNb12Y9y73ZMWLaDoYmo8HvS6mq/GypSnKmUSbzBkhzilbjjp55yjy/Dmgl2RapzmQdqVVlkz4NDrQ"
    "UDka+SJFUgQM2/gzdDQpcf7ZRZFeyoXaW+USaIQ2EchWyScnQaR7QTta2qI5oXC2mIlLXum41OFZq1wArTo5yF58sU1cc3eD"
    "xtO+ZnEWN1CTh9v3BnA4gE4bxl0C3ckklhVj6kri3aAXvaRn8zvC4K/NkO7UJRzWj/nKzV+U6NbTCko1c+H7WvoPQfeMrusr"
    "48Wb2mHrqSMxFaau8OqXVEfl/IF2VQd7M+hugGVh+lAzcyKur2BOe8OyML7mOsxVB881qLLdOTHcOYKZU+1PQfcc/50IyXBS"
    "BdRTYNwANEATiVSv+Q5NpokCQFcqCsoHpQ8Bjdtcd0BXov0ImCcBwKzbIbEDm9kz13OkatCJn+nluDr1HNDdaCI2cmor3alZ"
    "SkkmxJRDhhFoK6v3leu1mB/p6PyRyNY1OlLqJznwfwq6pzsszPI6ZG3JsOjeEWQd5L8fT46DTJ13uyh7CTS295vfKj0JdDdZ"
    "C3IHQm/ShcvUuzZ5t6O7jZKQFn0U/cyOzvFlTOm85E06ehBxzHB0bYPKMrQy77Nv39kyTUc7Qhzad+ZF0Os77wI98F/6YDaH"
    "v+3Ay9yIY+mIL8E0PcGo7MpC3Z8Jem3OQ697WzbN8MIlrPpa0K8YWRYusbswpYC+GfT1nAOEdwk0wPeCVpfXIafFKFa0Ty40"
    "ivpm0FdJuysPgb6iZcxacPZjQV8j7RSsC7S+kg3l1tINPhj0lamzRsH6hA+NF6waV0x/s28EbW8DvTrpPs66AFpqFwv1SuGn"
    "LZhnKpGybwNtgp11B+i10Cmv9AFpfs9cRMMFFubx6jhv5nhu9JtAu2hv3gJ6vjBKtejKybozqVFUSpSeRXBd8mMd/QC9CXTy"
    "hN8DegrPu+mqSOEJPl5QCIoI7vDU2mQRhOSyf6tEc+9/F+gArxfRjsninQXIOjPwOqXjulidUzvKc9bVFt998j4dHUlXL7xp"
    "twZipwwslMHeIqhMj4y1VgrZzvSHavE2JBdK23Hp9tT7XxMLF3NY8HwhgJ9bHUYc5/WLI2G2suJvlalm7uRlRqsFVqXPFyRc"
    "Khzd5CepO3nuXc7ubCfU2T8+n+yeRUFrwtDiPlv9cr4scAg0TgpfeotoLrznO3uZPwQrkf7XyqSYNiwtg3lhheRLpU8Kh3EL"
    "4EoSjzuX6EEuEIbceAdvFVtQ44Z48xOzcoGw9ha8nue65Lvy5oF+27LG7HrmMVdnOTY9fcHTWksMdL8f5K+8QSp6reB97+o0"
    "IVZq3wvaDEBvb32pJNtTi6DF/nobaRPP/lTQa8/+oNQl0NmrW39/0+l5+X3QRxfZgD6SCmRwWlgEqdOM6yjoY7liqPR9OAur"
    "IEy5TMciDAl0a3cUFxsHdH9r6/k6fkuQ01aDzrxKLNEx7bWbc6cLmYfiIyU9pvMUmY2QQGenhqFE+6U/3qmjG+/db3EmY7zJ"
    "jjVozrukfbgma1XuY2QKr/4kx4l1in+TT++yDq9o8//3L2wL30LOM2FICfOFuQCEM2GSaE7/lNnL4F9SxyXBv2WGyyJeSFl6"
    "eWv90b+kN/jNv9o/NCXo4CHTrDXCFx6YhzAAqYCHP+V95BDF1rGHzoaRrolPfHhruXdi8BoIOg4dUOYNpXe0Hhe3cDhGb3ov"
    "+K/1gNkzc4DOMhKzeWcaju+Ug95K0EXOl2fmnVGkEmiJgNdnckdM4BBe30xa3seNHwtabkhzxlEl0XLbziTQmhe6EhVmvHNa"
    "56DNlpYhAHaEai62BbzkpZeZ+ams3OtYPsbxmQDDx/6rTSEUPlySDLheVn0waBLDzbagZZ8vIKB31elv1j/C8qxvkIFOnWHs"
    "8Qwv+aOD4EqWFz//0rgkKY9OzuS1rxbRVZLDK0XkJCY8BQifDdqKFi1Bi2+dWFPo8OEfXwYdcmRQNKt86jAbINqCxN4EK6VS"
    "+IJLBYVcn0kqwYqaFTPyqX0uU9TWHw5adyVaLUp07Mh0StdLEu34CIYoawKFUtGWW5To2JmSgm/W0Xqso1mpGm9bnIEmEB3t"
    "TcICdHWmoKM5XzfX0Sr0v3k//NFWRwN6YHVgllygDw/BGLR0gVtcAegAfWJ1AOVpkMpsq8tIPBZ0GHDyU2wHdrQu7WjSKRQU"
    "DGELI9AszJl5vOWq49SONmHwYMP0DX3jW5R/a2To10ijOJrDOJO7GBn6DxOGgzxRz2D6KgHUeFwaL/ozxrP5AR9JVDD8FNY0"
    "iCND+UFGhiGgDmSADtVB6rNBn74cq7+r45CApWvAbH+2AqWz1oT+OAr0h4NWxeTAMsKSrSnSxvOq0B5Ua7xmnr3S35dHyrIQ"
    "Wvoj7jwXRpMc8NjuXRH9i7cjhhJcJfrOt799NWlJQeN3aYnlzrf9H+qce4EQNniRAAAAAElFTkSuQmCC"
)
_ICON_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAEAAAABACAMAAACdt4HsAAAAP1BMVEUAX3sARH4ARIEARYEABn8AQ34AAP8ARIEAS5AAQ34A"
    "HpgA//8ARYAAAAAARYIAPn4AAAAAAAAAAAAAAAAAAAAzBwJ3AAAAEHRSTlMT/GjKA2QBmxCGCgH+AC8JcrrwPgAAAWFJREFU"
    "eNrtlu1uwyAMRa/NZ7oA7/+2tSFJq25dCdF+TMJVG1zwwQbbCtaLggmYgAmYgAmYgAn4DKCVSL5tTM8TdNYDOn4OhUZCoJMh"
    "5AggypRhGYBN3ddwFiW84H4EIKlg5dSkiNWhZPPZA/jo4bNLtsB7FJ9Y7FGsKgkdIaTmQyntaZH4UKLQqA9Q0pNgH1iPbg/A"
    "vNuJ55x3lukCoAYbYvJ1fxZl2cB9gK9Yb49bJNiuB90Am2wdB7XZD26jdQGwbRoEJWZcFTcAMHoGfw2giwC66IHY8xggBDFG"
    "vRw+DyC5n5tkCvTfEYA4YHWJFj2+hWCLfHID2JsVWRrAVtlDKMmtkEzz+UVKemSvaQW0e9B6ShOWAndQr4q0sPIQRJ2GawCo"
    "EvW8aF3qTG6LY5QkBYS6LT0tUqnRoHZAHrKvnQGkBP61c7/p9KF2Fk0k487bb4l0pPKIPb0vpvl+8N8Ad7Not7N+LN5TAAAA"
    "AElFTkSuQmCC"
)


def logo_tag(white=False, height=50, extra=""):
    data = _LOGO_WHITE_B64 if white else _LOGO_BLUE_B64
    return ('<img src="data:image/png;base64,' + data + '" alt="Grupo MESS" '
            'style="height:' + str(height) + 'px;width:auto;' + extra + '"/>')


def page_icon():
    try:
        from PIL import Image
        return Image.open(io.BytesIO(base64.b64decode(_ICON_B64)))
    except Exception:
        return None


# ---------------------------------------------------------------------------
# CONSTANTES EDITABLES
# ---------------------------------------------------------------------------
MESS_BLUE = "#004580"
MESS_DARK = "#2c3e50"
ACCENT = "#ff6600"
RED = "#c0392b"

DEFAULT_MODEL = "gemini-3.8-flash"

FASES = {
    0: "0. Sin clasificar",
    1: "1. Propuesta",
    2: "2. Cotización",
    3: "3. Negociación",
    4: "4. Esperando PO",
}
# Probabilidad inicial por fase: calíbrala con tu histórico real de cierre.
PROB_DEFAULT = {1: 10, 2: 25, 3: 50, 4: 80}
# Antigüedad (días desde el registro) esperada como máximo por fase.
UMBRAL_DIAS = {0: 60, 1: 30, 2: 45, 3: 75, 4: 45}

SIGUIENTE_PASO = {
    0: "Revisar la etapa en SCOTT",
    1: "Validar necesidad y decisor; preparar cotización",
    2: "Confirmar recepción y criterios de decisión",
    3: "Resolver objeciones y fijar fecha de decisión",
    4: "Confirmar fecha de colocación de la PO",
}
OBJETIVO_FASE = {
    0: "Reactivar proyecto frío",
    1: "Calificar necesidad y decisor",
    2: "Dar seguimiento a cotización",
    3: "Manejar objeciones y fijar fecha de decisión",
    4: "Asegurar colocación de PO",
}
OBJETIVOS = [
    "Calificar necesidad y decisor",
    "Dar seguimiento a cotización",
    "Manejar objeciones y fijar fecha de decisión",
    "Asegurar colocación de PO",
    "Reactivar proyecto frío",
    "Ampliar la cuenta (cross-selling)",
    "Agendar visita técnica",
]
FOCO_MEDDPICC = {
    0: "validar si el proyecto sigue vivo y quién decide",
    1: "validar el dolor (Pain), al Champion y al comprador económico",
    2: "validar criterios de decisión (Decision Criteria) y competencia",
    3: "validar proceso de decisión (Decision Process), papeleo y comprador económico",
    4: "validar el proceso de compra interno y la fecha real de colocación de la PO",
}
DIAS_FRESCO = 25  # umbral de tu versión en GitHub: en tiempo (menos de 25 días) frente a rezagado
CIERRE_OPC = ["Todas", "Vencidas", "Próximos 7 días", "Este mes", "Próximo mes", "Rango personalizado"]

MK_TAGS = ("ESTRATEGIA", "LINKEDIN", "EMAILS", "WHATSAPP", "FICHA", "CALENDARIO", "KPIS")
MK_OBJETIVOS = [
    "Nutrir una cotización abierta",
    "Reactivar proyectos fríos",
    "Impulsar cross-selling de calibración o alta gama",
    "Asegurar la colocación de la PO",
    "Invitar a evento, webinar o capacitación",
    "Posicionar una marca o línea de equipos",
]
MK_CANALES = ["LinkedIn", "Correo", "WhatsApp", "Ficha o landing"]
MK_TONOS = ["Técnico y consultivo", "Ejecutivo y breve", "Cercano y humano"]
# Hechos públicos de Grupo MESS (mess.com.mx). Revísalos y ajústalos: la IA solo debe apoyarse en ellos.
MESS_HECHOS = [
    "Grupo MESS: 'all about metrology'. Laboratorio con acreditaciones ISO/IEC 17025 (confirma el alcance vigente antes de citar magnitudes o cifras).",
    "Servicios: calibración, inspección y medición, laboratorio de pruebas, equipo y consumibles, desarrollo profesional y desarrollo tecnológico.",
    "Alta gama: servicio a CMM, retrofit, ingeniería inversa, escáneres 3D, medición en sitio y renta de equipo.",
    "Marcas que representa o distribuye: Buehler, Wilson, Nikon, Kreon, Taylor Hobson, Zygo, Vicivision, Mark-10, Galdabini, Baty, Fluke, Powervar, Scanology.",
    "Capacitación técnica en línea (MESS Academy) y equipo seminuevo en renta o venta.",
    "Presencia en Querétaro, San Luis Potosí, Saltillo, Monterrey, Guadalajara y Aguascalientes.",
]
CANALES = ["WhatsApp", "Correo ejecutivo", "Llamada", "Visita presencial"]
PERFILES = ["Ingeniero de planta o calidad", "Comprador", "Director o dueño"]

# Cuentas que son el mismo grupo bajo distintos nombres en SCOTT.
# Revisa y ajusta esta lista: se infirió solo por el nombre del cliente.
ALIAS_CUENTAS = ["CNH", "SAFRAN", "ITP", "SIEMENS ENERGY", "ADIENT", "BROSE", "DANA", "STEERINGMEX"]

# ---------------------------------------------------------------------------
# PÁGINA Y ESTILOS
# ---------------------------------------------------------------------------
st.set_page_config(page_title="SAIV | Radar Comercial MESS", page_icon=page_icon(), layout="wide", initial_sidebar_state="expanded")

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Montserrat:wght@400;500;600;700;800&display=swap');
html, body, [class*="css"], .stApp, button, input, textarea, select {
    font-family: 'Montserrat', sans-serif !important;
}
#MainMenu {visibility: hidden;}
footer {visibility: hidden;}
.stApp { background-color: #f4f7fb; }
h1, h2, h3, h4 { color: #004580; font-weight: 700; }
.header-mess {
    display: flex; align-items: center; gap: 26px;
    background: linear-gradient(135deg, #004580 0%, #002b52 100%);
    padding: 18px 28px; border-radius: 10px; margin-bottom: 18px;
    border-bottom: 4px solid #ff6600; box-shadow: 0 4px 10px rgba(0,43,82,0.18);
}
.header-text { border-left: 1px solid rgba(255,255,255,0.35); padding-left: 24px; }
.header-text h1 { color: #ffffff; font-size: 22px; font-weight: 800; margin: 0; text-transform: uppercase; letter-spacing: 0.5px; }
.header-text p { font-size: 12px; margin: 4px 0 0 0; color: #cbd5e1; text-transform: uppercase; font-weight: 600; letter-spacing: 1px; }
.kpi-card {
    background: #ffffff; border: 1px solid #dde6f0; border-left: 6px solid #004580;
    border-radius: 8px; padding: 12px 16px; box-shadow: 0 2px 6px rgba(0,69,128,0.08); margin-bottom: 10px;
}
.kpi-label { font-size: 0.7rem; text-transform: uppercase; letter-spacing: 0.05em; color: #2c3e50; font-weight: 600; }
.kpi-value { font-size: 1.5rem; font-weight: 700; color: #004580; line-height: 1.2; }
.kpi-sub { font-size: 0.72rem; color: #64748b; }
section[data-testid="stSidebar"] { background-color: #ffffff; border-right: 1px solid #dde6f0; }
.sidebar-brand { padding: 4px 0 14px 0; margin-bottom: 8px; border-bottom: 3px solid #004580; }
div[data-baseweb="tab-list"] button[aria-selected="true"] { color: #004580 !important; font-weight: 700; }
div[data-baseweb="tab-highlight"] { background-color: #ff6600 !important; }
span[data-baseweb="tag"] { background-color: #004580 !important; color: #ffffff !important; }
.stButton > button { background-color: #004580; color: #ffffff; border: none; border-radius: 6px; font-weight: 600; }
.stButton > button:hover { background-color: #00335f; color: #ffffff; }
.login-brand { text-align: center; padding: 30px 0 10px 0; }
.brand-footer { text-align: center; color: #64748b; font-size: 0.72rem; margin-top: 30px; padding-top: 14px; border-top: 1px solid #dde6f0; }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)


# ---------------------------------------------------------------------------
# SEGURIDAD
# ---------------------------------------------------------------------------
def check_password():
    if st.session_state.get("auth_ok"):
        return True
    st.sidebar.markdown("### Acceso restringido")
    try:
        secret = str(st.secrets["mi_contrasena"])
    except Exception:
        secret = ""
    if not secret:
        st.sidebar.error("Falta 'mi_contrasena' en st.secrets. La app permanece bloqueada.")
        return False
    pwd = st.sidebar.text_input("Contraseña", type="password", key="pwd_input")
    if pwd:
        if hmac.compare_digest(pwd.encode("utf-8"), secret.encode("utf-8")):
            st.session_state["auth_ok"] = True
            st.rerun()
        else:
            st.sidebar.error("Contraseña incorrecta.")
    return False


if not check_password():
    st.markdown('<div class="login-brand">' + logo_tag(False, 90) + "</div>", unsafe_allow_html=True)
    st.info("Ingresa tu contraseña en el menú lateral para continuar.")
    st.stop()

st.markdown(
    '<div class="header-mess">' + logo_tag(True, 56)
    + '<div class="header-text"><h1>Radar Comercial</h1>'
    "<p>SAIV | Sistema Automatizado de Ingeniería de Ventas y Copiloto KAM</p></div></div>",
    unsafe_allow_html=True,
)


# ===========================================================================
# LÓGICA DE DATOS (inicio)
# ===========================================================================
def norm(text):
    s = "" if text is None else str(text)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.upper().strip()


def decode_bytes(raw):
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def detect_delimiter(text):
    try:
        return csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
    except Exception:
        return ","


# Reparaciones seguras de letras perdidas en SCOTT (el '?' ya viene así en el CSV).
# Solo patrones inequívocos: no se adivina una letra para cualquier '?'.
REPARACIONES = [
    ("ci?n", "ción"), ("CI?N", "CIÓN"), ("Ci?n", "Ción"),
    ("?xico", "éxico"), ("?XICO", "ÉXICO"),
    ("aut?nom", "autónom"), ("AUT?NOM", "AUTÓNOM"), ("Aut?nom", "Autónom"),
    ("se?o", "seño"), ("SE?O", "SEÑO"), ("Se?o", "Seño"),
    ("a?o", "año"), ("A?O", "AÑO"),
    ("metr?a", "metría"), ("METR?A", "METRÍA"), ("Metr?a", "Metría"),
    ("?ptic", "óptic"), ("?PTIC", "ÓPTIC"), ("?ptic", "Óptic"),
]


def reparar_texto(x):
    t = "" if x is None else str(x)
    for a, b in REPARACIONES:
        t = t.replace(a, b)
    return re.sub(r"^\?+(?=[A-Za-zÁÉÍÓÚÑáéíóúñ])", "", t).strip()


def parse_valor(v):
    """'18,900.00 MXN' -> (18900.0, 'MXN'). Sin moneda explícita se toma MXN."""
    s = "" if v is None else str(v)
    m = re.search(r"-?\d[\d,]*\.?\d*", s)
    try:
        amt = float(m.group().replace(",", "")) if m else 0.0
    except ValueError:
        amt = 0.0
    up = s.upper()
    cur = "USD" if ("USD" in up or "US$" in up) else "MXN"
    return amt, cur


def classify_fase(etapa):
    t = norm(etapa)
    if re.search(r"ESPERANDO|\bPO\b|ORDEN DE COMPRA", t):
        return 4
    if "NEGOCI" in t:
        return 3
    if "COTIZ" in t:
        return 2
    if "PROPUESTA" in t:
        return 1
    return 0


def tipo_de_area(area):
    a = norm(area)
    if a.startswith("VENTAS"):
        return "Equipos y consumibles"
    if a.startswith(("LABORATORIO", "SERVICIOS")) or "DIMENSIONAL" in a or "CALIBRAC" in a:
        return "Servicios y laboratorio"
    return "Otros (revisar)"


def cuenta_de(cliente):
    n = norm(cliente)
    for key in ALIAS_CUENTAS:
        if n == key or n.startswith(key + " "):
            return key
    return n


@st.cache_data(show_spinner=False)
def load_pipeline(raw_bytes, hoy_iso):
    """Lee el CSV de SCOTT. Devuelve (DataFrame, lista de avisos)."""
    avisos = []
    hoy = pd.Timestamp(hoy_iso)
    text = decode_bytes(raw_bytes)
    rows = list(csv.reader(io.StringIO(text), delimiter=detect_delimiter(text)))

    h = None
    for i, r in enumerate(rows):
        joined = " ".join(norm(c) for c in r)
        if "PROYECTO" in joined and "CLIENTE" in joined:
            h = i
            break
    if h is None:
        return None, ["No se encontró la fila de encabezados con PROYECTO y CLIENTE."]

    header = [norm(c) for c in rows[h]]

    def find(*keys):
        for k in keys:
            for pos, n in enumerate(header):
                if k in n:
                    return pos
        return None

    pos = {
        "proy": find("PROYECTO", "FOLIO"),
        "reg": find("FECHA DE REGISTRO"),
        "desc": find("DESCRIPCION"),
        "cli": find("CLIENTE"),
        "area": find("AREA"),
        "cierre": find("FECHA DE CIERRE"),
        "valor": find("VALOR", "MONTO", "IMPORTE"),
        "cot": find("COTIZACION"),
        "etapa": find("ETAPA", "FASE"),
        "estatus": find("ESTATUS", "STATUS"),
    }
    if pos["proy"] is None or pos["cli"] is None:
        return None, ["No se pudieron identificar las columnas PROYECTO y CLIENTE."]
    if pos["valor"] is None:
        avisos.append("No se encontró la columna VALOR; los montos quedan en cero.")
    if pos["etapa"] is None:
        avisos.append("No se encontró la columna ETAPA; todo queda sin clasificar.")

    def cell(r, key):
        p = pos.get(key)
        if p is None or p >= len(r):
            return ""
        return str(r[p]).strip()

    recs = []
    cur = None
    for r in rows[h + 1:]:
        if not any(str(c).strip() for c in r):
            continue
        folio = cell(r, "proy")
        if folio:
            cur = {
                "Folio": folio,
                "Registro_txt": cell(r, "reg"),
                "Descripcion": cell(r, "desc"),
                "Cliente": cell(r, "cli"),
                "Razon_Social": "",
                "Area": cell(r, "area"),
                "Cierre_txt": cell(r, "cierre"),
                "Valor_txt": cell(r, "valor"),
                "Cotizacion": cell(r, "cot"),
                "Etapa": cell(r, "etapa"),
                "Estatus": cell(r, "estatus"),
            }
            recs.append(cur)
        elif cur is not None:
            rs = cell(r, "cli")
            if rs and not cur["Razon_Social"]:
                cur["Razon_Social"] = rs

    if not recs:
        return None, ["El archivo no contiene proyectos debajo de los encabezados."]

    d = pd.DataFrame(recs)
    d = d.drop_duplicates(subset="Folio", keep="first").reset_index(drop=True)
    for c in ("Descripcion", "Cliente", "Razon_Social", "Area"):
        d[c] = d[c].map(reparar_texto)

    sin_cli = d["Cliente"].map(lambda x: norm(x) in ("", "DIRECCION FISCAL"))
    if sin_cli.any():
        d.loc[sin_cli, "Cliente"] = d.loc[sin_cli, "Razon_Social"].replace("", "SIN CLIENTE")
        avisos.append(str(int(sin_cli.sum())) + " proyecto(s) sin nombre de cliente válido; se usó la razón social.")

    vals = d["Valor_txt"].map(parse_valor)
    d["Monto_USD"] = [a if c == "USD" else 0.0 for a, c in vals]
    d["Monto_MXN"] = [a if c == "MXN" else 0.0 for a, c in vals]
    d["Sin_Cotizar"] = (d["Monto_USD"] == 0) & (d["Monto_MXN"] == 0)

    d["Cotizacion"] = d["Cotizacion"].map(lambda x: "" if norm(x) in ("", "SIN REGISTRO") else x)
    d["Fase_Num"] = d["Etapa"].map(classify_fase)
    d["Fase"] = d["Fase_Num"].map(FASES)
    n_sin_clasif = int((d["Fase_Num"] == 0).sum())
    if n_sin_clasif:
        avisos.append(str(n_sin_clasif) + " proyecto(s) con etapa no reconocida quedaron como 'Sin clasificar'.")

    d["Tipo"] = d["Area"].map(tipo_de_area)
    d["Cuenta"] = d["Cliente"].map(cuenta_de)

    d["Registro"] = pd.to_datetime(d["Registro_txt"], dayfirst=True, errors="coerce")
    d["Cierre_Est"] = pd.to_datetime(d["Cierre_txt"], dayfirst=True, errors="coerce")
    n_sin_fecha = int(d["Registro"].isna().sum())
    if n_sin_fecha:
        avisos.append(str(n_sin_fecha) + " proyecto(s) sin fecha de registro válida; su antigüedad se toma como 0.")
    d["Dias_Registro"] = (hoy - d["Registro"].dt.normalize()).dt.days.fillna(0).astype(int)
    d["Dias_a_Cierre"] = (d["Cierre_Est"].dt.normalize() - hoy).dt.days

    return d, avisos


def enrich(df, tc, probs, rates):
    """Agrega equivalencias, ponderado, comisión estimada y alertas."""
    d = df.copy()
    d["Total_USD_Eq"] = d["Monto_USD"] + d["Monto_MXN"] / tc
    d["Prob"] = d["Fase_Num"].map(lambda f: probs.get(int(f), 10) / 100.0)
    d["Ponderado"] = d["Total_USD_Eq"] * d["Prob"]
    d["Tasa_Com"] = d["Tipo"].map(lambda t: rates.get(t, 0.0) / 100.0)
    d["Comision_Est"] = d["Ponderado"] * d["Tasa_Com"]

    alertas, semaforos = [], []
    for _, r in d.iterrows():
        flags = []
        if r["Sin_Cotizar"]:
            flags.append("Cotizar")
        dc = r["Dias_a_Cierre"]
        if pd.notna(dc) and dc < 0:
            flags.append("Cierre vencido")
        u = UMBRAL_DIAS.get(int(r["Fase_Num"]), 60)
        if r["Dias_Registro"] > 2 * u:
            flags.append("Depurar")
        elif r["Dias_Registro"] > u:
            flags.append("Dar seguimiento")
        alertas.append(", ".join(flags))
        if "Cierre vencido" in flags or "Depurar" in flags:
            semaforos.append("Rojo")
        elif flags:
            semaforos.append("Ámbar")
        else:
            semaforos.append("Verde")
    d["Alerta"] = alertas
    d["Semaforo"] = semaforos
    d["Siguiente_Paso"] = [
        "Cotizar y enviar al cliente" if s else SIGUIENTE_PASO.get(int(f), "")
        for s, f in zip(d["Sin_Cotizar"], d["Fase_Num"])
    ]
    return d


def filtrar_cierre(d, modo, rango, hoy):
    """Filtra por la fecha de cierre estimada. Sin fecha solo aparece en 'Todas'."""
    c = d["Cierre_Est"]
    if modo == "Vencidas":
        return d[c < hoy]
    if modo == "Próximos 7 días":
        return d[(c >= hoy) & (c <= hoy + pd.Timedelta(days=7))]
    if modo == "Este mes":
        return d[c.dt.to_period("M") == hoy.to_period("M")]
    if modo == "Próximo mes":
        return d[c.dt.to_period("M") == (hoy.to_period("M") + 1)]
    if modo == "Rango personalizado" and rango is not None and len(rango) == 2:
        ini, fin = pd.Timestamp(rango[0]), pd.Timestamp(rango[1])
        return d[(c >= ini) & (c <= fin)]
    return d


# ===========================================================================
# LÓGICA DE DATOS (fin)
# ===========================================================================


def fmt_usd(x):
    try:
        return "$" + format(float(x), ",.0f")
    except Exception:
        return "$0"


def show_chart(chart):
    try:
        st.altair_chart(chart, width="stretch")
    except TypeError:
        st.altair_chart(chart, use_container_width=True)


def kpi_card(label, value, sub=""):
    st.markdown(
        '<div class="kpi-card"><div class="kpi-label">' + label + "</div>"
        '<div class="kpi-value">' + value + "</div>"
        '<div class="kpi-sub">' + sub + "</div></div>",
        unsafe_allow_html=True,
    )


DISPLAY = {
    "Monto_USD": "USD",
    "Monto_MXN": "MXN",
    "Total_USD_Eq": "Total Eq. USD",
    "Ponderado": "Ponderado USD",
    "Comision_Est": "Comisión est. USD",
    "Dias_Registro": "Días reg.",
    "Dias_a_Cierre": "Días a cierre",
    "Cierre_Est": "Cierre est.",
    "Siguiente_Paso": "Siguiente paso",
    "Semaforo": "Semáforo",
    "Descripcion": "Descripción",
    "Cotizacion": "Cotización",
    "Fase": "Fase",
}
SEMAFORO_CSS = {
    "Rojo": "color:#c0392b;font-weight:700",
    "Ámbar": "color:#b9770e;font-weight:700",
    "Verde": "color:#1e8449;font-weight:700",
}


def show_table(d, cols, key=None):
    if d is None or d.empty:
        st.info("Sin registros para esta vista.")
        return
    t = d[cols].copy()
    if "Descripcion" in t.columns:
        t["Descripcion"] = t["Descripcion"].astype(str).str.slice(0, 70)
    if "Cierre_Est" in t.columns:
        t["Cierre_Est"] = t["Cierre_Est"].dt.strftime("%d/%m/%Y").fillna("")
    t = t.rename(columns=DISPLAY)
    fmt = {}
    for c in ("USD", "MXN"):
        if c in t.columns:
            fmt[c] = "${:,.2f}"
    for c in ("Total Eq. USD", "Ponderado USD", "Comisión est. USD"):
        if c in t.columns:
            fmt[c] = "${:,.0f}"
    sty = t.style.format(fmt, na_rep="")
    if "Semáforo" in t.columns:
        try:
            sty = sty.map(lambda v: SEMAFORO_CSS.get(v, ""), subset=["Semáforo"])
        except AttributeError:
            sty = sty.applymap(lambda v: SEMAFORO_CSS.get(v, ""), subset=["Semáforo"])
    try:
        st.dataframe(sty, hide_index=True, width="stretch")
    except TypeError:
        st.dataframe(sty, hide_index=True, use_container_width=True)


def totals_caption(d):
    if d is None or d.empty:
        return
    st.caption(
        "Totales de esta vista: " + str(len(d)) + " proyectos | USD "
        + format(d["Monto_USD"].sum(), ",.2f") + " | MXN " + format(d["Monto_MXN"].sum(), ",.2f")
        + " | Total equivalente USD " + format(d["Total_USD_Eq"].sum(), ",.0f")
    )


# ---------------------------------------------------------------------------
# IA (GEMINI)
# ---------------------------------------------------------------------------
def call_gemini(prompt):
    if genai is None:
        return None, "Falta instalar la librería google-generativeai."
    key = None
    model_name = DEFAULT_MODEL
    try:
        key = st.secrets.get("gemini_api_key") or st.secrets.get("GEMINI_API_KEY")
        model_name = st.secrets.get("gemini_model") or DEFAULT_MODEL
    except Exception:
        pass
    if not key:
        return None, "Falta 'gemini_api_key' en st.secrets."
    try:
        genai.configure(api_key=key)
        resp = genai.GenerativeModel(model_name).generate_content(prompt)
        return resp.text, None
    except Exception as e:
        return None, "Error al consultar Gemini (" + model_name + "): " + str(e)


def parse_tags(text, tags=("ANALYSIS", "MESSAGE", "OBJECTIONS", "LOG")):
    out = {}
    for tag in tags:
        pat = "<" + tag + ">(.*?)(?=</" + tag + ">|<[A-Z]+>|$)"
        m = re.search(pat, text or "", re.S | re.I)
        out[tag] = m.group(1).strip() if m else ""
    return out


def build_prompt(ctx):
    montos = []
    if ctx["usd"] > 0:
        montos.append(format(ctx["usd"], ",.2f") + " USD")
    if ctx["mxn"] > 0:
        montos.append(format(ctx["mxn"], ",.2f") + " MXN")
    monto_txt = " y ".join(montos) if montos else "Aún sin cotizar (sin monto): no menciones precios en el mensaje."
    lines = [
        "Actúa como Javier Camacho, Key Account Manager (KAM) B2B de MESS Servicios Metrológicos en Querétaro, México.",
        "Eres experto en metrología dimensional, calibración acreditada ISO/IEC 17025 y equipos de alta gama (Zeiss, Taylor Hobson, Wilson, Kreon).",
        "Hablas como asesor técnico de metrología industrial: trazabilidad, incertidumbre, repetibilidad, paros de línea, retrabajo y certeza dimensional en piso de producción.",
        "Tono consultivo, humano y directo, sin exageraciones técnicas ni lenguaje de marketing inflado. Frases claras y cortas.",
        "No inventes precios, fechas de entrega, certificaciones, referencias ni datos que no estén en el contexto.",
        "",
        "DATOS DE LA OPORTUNIDAD:",
        "- Cuenta: " + ctx["cuenta"],
        "- Folios agrupados: " + ctx["folios"],
        "- Descripción de equipos o servicios: " + ctx["descripcion"],
        "- Monto del paquete: " + monto_txt,
        "- Fase actual: " + ctx["fase"],
        "- Días desde el registro: " + str(ctx["dias"]),
        "- Cierre estimado en SCOTT: " + ctx["cierre"],
        "- Cotizaciones: " + (ctx["cotizaciones"] or "Sin número de cotización registrado"),
        "- Propuestas a incluir:",
        ctx["desglose"],
        "- Objetivo de esta acción: " + ctx["objetivo"],
        "- Canal: " + ctx["canal"],
        "- Interlocutor: " + ctx["perfil"],
        "- Giro industrial: " + (ctx["giro"] or "No especificado"),
        "- Contexto adicional: " + (ctx["contexto"] or "Ninguno"),
        "",
        "METODOLOGÍA (aplícala de forma invisible y natural):",
        "- SPIN Selling: preguntas de Situación, Problema, Implicación y Necesidad-Solución sobre el riesgo de medir mal o calibrar tarde.",
        "- MEDDPICC: en esta fase el foco es " + ctx["foco"] + ". Señala qué elementos faltan por validar y qué pregunta los destapa.",
        "- Sandler: dale al cliente permiso de decir que no, sin presión y sin forzar una decisión. Busca una respuesta honesta, no una promesa.",
        "",
        "FORMATO: responde ÚNICAMENTE con estas cuatro etiquetas XML, en este orden, sin texto fuera de ellas y sin bloques de código:",
        "<ANALYSIS> Diagnóstico breve: qué sabemos, qué falta validar (lista corta MEDDPICC) y la mejor jugada para esta fase. </ANALYSIS>",
        "<MESSAGE> Dos versiones listas para enviar por el canal indicado, separadas por una línea con tres guiones (---): primero la versión BREVE y después la COMPLETA. Incluye al menos una pregunta SPIN y un cierre con permiso de decir no. </MESSAGE>",
        "REGLAS PARA <MESSAGE>: el texto se mostrará en una caja de texto plano. PROHIBIDO usar tablas Markdown (nada de barras verticales ni :---) y PROHIBIDO usar asteriscos.",
        "En la versión COMPLETA presenta cada propuesta en una lista vertical, con una línea en blanco entre propuestas, con este formato exacto:",
        "• Cotización [Número] - [Equipo o servicio]",
        "Impacto en planta: [beneficio técnico metrológico o cumplimiento normativo IATF/ISO, sin inventar certificaciones adicionales]",
        "Inversión: [monto exacto en USD o MXN; si no hay monto escribe 'Por cotizar']",
        "La versión BREVE no lleva la lista: solo menciona el paquete y una pregunta.",
        "<OBJECTIONS> Las dos o tres objeciones más probables (precio, tiempos, prioridad interna) y cómo responderlas con enfoque Sandler. </OBJECTIONS>",
        "<LOG> Resumen de dos a cuatro líneas para registrar en SCOTT: acción, interlocutor y siguiente paso. </LOG>",
    ]
    return "\n".join(lines)


def clean_message(txt):
    """Texto plano: quita asteriscos y líneas de tablas Markdown sin borrar contenido útil."""
    out = []
    for line in str(txt or "").replace("**", "").splitlines():
        if re.match(r"^\s*\|.*\|\s*$", line) or re.match(r"^\s*:?-{3,}:?(\s*\|\s*:?-{3,}:?)+\s*$", line):
            continue
        out.append(line.rstrip())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip()


def whatsapp_link(message, phone=""):
    digits = re.sub(r"\D", "", phone or "")
    base_url = "https://wa.me/" + digits if digits else "https://wa.me/"
    return base_url + "?text=" + quote(message or "")


def gcal_link(title, details, start_dt, minutes=30):
    end_dt = start_dt + timedelta(minutes=minutes)
    f = "%Y%m%dT%H%M%S"
    return (
        "https://calendar.google.com/calendar/render?action=TEMPLATE&text=" + quote(title)
        + "&details=" + quote((details or "")[:1500])
        + "&dates=" + start_dt.strftime(f) + "/" + end_dt.strftime(f)
    )


def build_docx(ctx, tags, mensaje, log_txt):
    doc = Document()
    doc.add_heading("Estrategia KAM | MESS Servicios Metrológicos", 0)
    doc.add_paragraph("Cuenta: " + ctx["cuenta"] + "\nFolios: " + ctx["folios"]
                      + "\nFecha: " + datetime.now().strftime("%d/%m/%Y"))
    doc.add_paragraph("Fase: " + ctx["fase"] + " | Objetivo: " + ctx["objetivo"] + " | Canal: " + ctx["canal"])
    doc.add_paragraph("USD: " + format(ctx["usd"], ",.2f") + " | MXN: " + format(ctx["mxn"], ",.2f"))
    doc.add_paragraph("Cotizaciones: " + (ctx["cotizaciones"] or "N/D") + "\n" + ctx["desglose"])
    for titulo, txt in [
        ("Análisis", tags.get("ANALYSIS", "")),
        ("Mensaje", mensaje),
        ("Manejo de objeciones", tags.get("OBJECTIONS", "")),
        ("Registro para SCOTT", log_txt),
    ]:
        doc.add_heading(titulo, level=2)
        doc.add_paragraph(txt or "Sin contenido.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# SIDEBAR: DATOS, PARÁMETROS Y FILTROS
# ---------------------------------------------------------------------------
with st.sidebar:
    st.markdown('<div class="sidebar-brand">' + logo_tag(False, 54) + "</div>", unsafe_allow_html=True)
    st.markdown("### Carga de datos")
    archivo = st.file_uploader("Extracción CSV de SCOTT", type=["csv"])

    st.markdown("### Parámetros")
    tc = st.number_input("Tipo de cambio (MXN por USD)", min_value=1.0, value=18.0, step=0.5)
    meta = st.number_input("Meta mensual (USD)", min_value=0.0, value=80000.0, step=5000.0)
    cerrado = st.number_input(
        "Cerrado del mes (USD, de tu otra app)", min_value=0.0, value=0.0, step=1000.0,
        help="Los ganados no vienen en este CSV; captúralo para calcular faltante y cobertura.",
    )
    with st.expander("Probabilidad por fase (%)"):
        probs = {
            1: st.number_input("Propuesta", 0, 100, PROB_DEFAULT[1], step=5),
            2: st.number_input("Cotización", 0, 100, PROB_DEFAULT[2], step=5),
            3: st.number_input("Negociación", 0, 100, PROB_DEFAULT[3], step=5),
            4: st.number_input("Esperando PO", 0, 100, PROB_DEFAULT[4], step=5),
        }
    with st.expander("Comisión estimada (opcional, %)"):
        st.caption("Captura tu tasa efectiva por tipo. Con cero no se muestra la columna.")
        rates = {
            "Equipos y consumibles": st.number_input("Equipos y consumibles", 0.0, 100.0, 0.0, step=0.5),
            "Servicios y laboratorio": st.number_input("Servicios y laboratorio", 0.0, 100.0, 0.0, step=0.5),
            "Otros (revisar)": st.number_input("Otros", 0.0, 100.0, 0.0, step=0.5),
        }

hoy = pd.Timestamp.today().normalize()
base = None
avisos = []
if archivo is not None:
    try:
        raw_df, avisos = load_pipeline(archivo.getvalue(), hoy.strftime("%Y-%m-%d"))
        if raw_df is not None:
            base = enrich(raw_df, tc, probs, rates)
        else:
            st.sidebar.error("; ".join(avisos))
    except Exception as e:
        st.sidebar.error("No se pudo procesar el CSV: " + str(e))

df = pd.DataFrame()
if base is not None:
    with st.sidebar:
        st.markdown("### Búsqueda y filtros")
        f_txt = st.text_input("Buscar", placeholder="Cliente, folio, equipo, marca, cotización...")
        f_fase = st.multiselect("Fase", [FASES[i] for i in sorted(base["Fase_Num"].unique())])
        f_tipo = st.multiselect("Tipo", sorted(base["Tipo"].unique()))
        f_area = st.multiselect("Área o marca", sorted(base["Area"].unique()))
        f_cuenta = st.multiselect("Cuenta", sorted(base["Cuenta"].unique()))
        f_min = st.number_input("Monto mínimo (USD eq.)", min_value=0, value=0, step=1000)
        f_dias = st.number_input("Antigüedad mínima (días)", min_value=0, value=0, step=15)
        f_frescos = st.checkbox("Solo en tiempo (menos de " + str(DIAS_FRESCO) + " días)")
        f_rezag = st.checkbox("Solo rezagados (" + str(DIAS_FRESCO) + " días o más)")
        f_sin_cot = st.checkbox("Solo sin cotizar")
        f_riesgo = st.checkbox("Solo en rojo (cierre vencido o por depurar)")
        st.markdown("##### Fecha de cierre")
        f_cierre = st.selectbox("Cierre estimado", CIERRE_OPC)
        f_rango = None
        if f_cierre == "Rango personalizado":
            f_rango = st.date_input("Del - al", value=(hoy.date(), (hoy + pd.Timedelta(days=30)).date()))
            if len(f_rango) != 2:
                st.caption("Elige la fecha final para aplicar el rango.")
        _cmin, _cmax = base["Cierre_Est"].min(), base["Cierre_Est"].max()
        if pd.notna(_cmin):
            st.caption("Cierres en SCOTT: " + _cmin.strftime("%d/%m/%Y") + " a " + _cmax.strftime("%d/%m/%Y"))

    df = base.copy()
    if f_txt:
        needle = f_txt.strip().lower()
        hay = (
            df["Cliente"].astype(str) + " " + df["Razon_Social"].astype(str) + " " + df["Folio"].astype(str)
            + " " + df["Descripcion"].astype(str) + " " + df["Area"].astype(str) + " " + df["Cotizacion"].astype(str)
        ).str.lower()
        df = df[hay.str.contains(needle, regex=False)]
    if f_fase:
        df = df[df["Fase"].isin(f_fase)]
    if f_tipo:
        df = df[df["Tipo"].isin(f_tipo)]
    if f_area:
        df = df[df["Area"].isin(f_area)]
    if f_cuenta:
        df = df[df["Cuenta"].isin(f_cuenta)]
    if f_min > 0:
        df = df[df["Total_USD_Eq"] >= f_min]
    if f_dias > 0:
        df = df[df["Dias_Registro"] >= f_dias]
    if f_frescos:
        df = df[df["Dias_Registro"] < DIAS_FRESCO]
    if f_rezag:
        df = df[df["Dias_Registro"] >= DIAS_FRESCO]
    if f_sin_cot:
        df = df[df["Sin_Cotizar"]]
    if f_riesgo:
        df = df[df["Semaforo"] == "Rojo"]
    df = filtrar_cierre(df, f_cierre, f_rango, hoy)
    st.sidebar.caption("Mostrando " + str(len(df)) + " de " + str(len(base)) + " proyectos")

hay_comision = any(v > 0 for v in rates.values())

tab_dash, tab_lab, tab_mk = st.tabs(["Dashboard de Oportunidades", "Laboratorio IA", "Marketing Digital"])


# ---------------------------------------------------------------------------
# DASHBOARD
# ---------------------------------------------------------------------------
def render_dashboard():
    if base is None:
        st.info("Sube la extracción CSV de SCOTT en el panel lateral para ver tu pipeline.")
        return
    for a in avisos:
        st.warning(a)
    if df.empty:
        st.info("Ningún proyecto coincide con los filtros activos.")
        return

    # --- KPIs ---------------------------------------------------------
    n_usd = int((df["Monto_USD"] > 0).sum())
    n_mxn = int((df["Monto_MXN"] > 0).sum())
    c1, c2, c3 = st.columns(3)
    with c1:
        kpi_card("Pipeline en USD", "$" + format(df["Monto_USD"].sum(), ",.2f"), str(n_usd) + " proyectos con monto en USD")
    with c2:
        kpi_card("Pipeline en MXN", "$" + format(df["Monto_MXN"].sum(), ",.2f"), str(n_mxn) + " proyectos con monto en MXN")
    with c3:
        kpi_card("Total equivalente USD", fmt_usd(df["Total_USD_Eq"].sum()), "Tipo de cambio " + str(tc))

    fin_mes = hoy + pd.offsets.MonthEnd(0)
    mes = df[df["Cierre_Est"] <= fin_mes]
    ponderado = float(df["Ponderado"].sum())
    faltante = max(0.0, meta - cerrado)
    cobertura = (ponderado / faltante) if faltante > 0 else None

    k1, k2, k3, k4 = st.columns(4)
    with k1:
        kpi_card("Probable (ponderado)", fmt_usd(ponderado), "Monto por probabilidad de fase")
    with k2:
        kpi_card("Cierre estimado este mes", fmt_usd(mes["Total_USD_Eq"].sum()),
                 str(len(mes)) + " proyectos | ponderado " + fmt_usd(mes["Ponderado"].sum()))
    with k3:
        kpi_card("Faltante para la meta", fmt_usd(faltante), "Meta " + fmt_usd(meta) + " menos cerrado " + fmt_usd(cerrado))
    with k4:
        kpi_card("Cobertura de pipeline", (format(cobertura, ".1f") + "x") if cobertura is not None else "Meta cubierta",
                 "Probable entre faltante")

    top_date = df["Cierre_Est"].value_counts()
    if len(top_date) and top_date.iloc[0] / len(df) >= 0.30:
        st.warning(
            str(round(top_date.iloc[0] / len(df) * 100)) + "% de los proyectos comparten la fecha de cierre estimada "
            + top_date.index[0].strftime("%d/%m/%Y")
            + ". Si es una fecha por defecto de SCOTT, el cierre estimado del mes está sobrestimado."
        )
    st.caption("La antigüedad se mide desde la fecha de registro; el CSV no trae fecha del último movimiento.")

    # --- Vista de moneda ------------------------------------------------
    vista = st.radio("Vista de gráficas", ["Global (USD equivalente)", "Solo USD", "Solo MXN"], horizontal=True)
    if vista.startswith("Global"):
        dv, col, cur = df, "Total_USD_Eq", "USD eq."
    elif vista == "Solo USD":
        dv, col, cur = df[df["Monto_USD"] > 0], "Monto_USD", "USD"
    else:
        dv, col, cur = df[df["Monto_MXN"] > 0], "Monto_MXN", "MXN"
    dv = dv.assign(Valor_Vista=dv[col], Pond_Vista=dv[col] * dv["Prob"])

    # --- Dónde actuar hoy ----------------------------------------------
    st.markdown("### Dónde actuar hoy")
    cols_base = ["Folio", "Cliente", "Descripcion", "Fase", "Monto_USD", "Monto_MXN", "Total_USD_Eq",
                 "Ponderado", "Cierre_Est", "Dias_Registro", "Alerta", "Siguiente_Paso"]
    cols_top = list(cols_base)
    orden = "Ponderado"
    if hay_comision:
        cols_top.insert(8, "Comision_Est")
        orden = "Comision_Est"
    t1, t2, t3, t4 = st.tabs([
        "Top 10 valor esperado", "Cierre próximo o vencido", "Por cotizar", "Por depurar o rescatar",
    ])
    with t1:
        st.caption("Ordenado por " + ("comisión estimada." if hay_comision else "monto ponderado por probabilidad de fase."))
        show_table(df[df["Total_USD_Eq"] > 0].sort_values(orden, ascending=False).head(10), cols_top)
    with t2:
        prox = df[df["Dias_a_Cierre"] <= 7].sort_values("Dias_a_Cierre")
        st.caption("Proyectos con cierre estimado en los próximos 7 días o ya vencido.")
        show_table(prox.head(25), cols_base)
    with t3:
        pend = df[df["Sin_Cotizar"]].sort_values("Dias_Registro", ascending=False)
        st.caption("Proyectos con valor 0.00: aún no hay monto cotizado. Sin cotización no hay cierre.")
        show_table(pend.head(30), ["Folio", "Cliente", "Descripcion", "Fase", "Area", "Dias_Registro", "Cierre_Est", "Siguiente_Paso"])
    with t4:
        rojo = df[df["Semaforo"] == "Rojo"].sort_values("Total_USD_Eq", ascending=False)
        st.caption("Cierre vencido o antigüedad mayor al doble de lo esperado para su fase. Candidatos a rescate (cierre negativo de Sandler) o a cerrar en SCOTT.")
        show_table(rojo.head(30), cols_base)

    # --- Gráficas -----------------------------------------------------------
    st.markdown("### Concentración y embudo (" + cur + ")")
    g1, g2 = st.columns(2)
    with g1:
        st.markdown("**Pareto 80/20 por cuenta**")
        par = dv.groupby("Cuenta")["Valor_Vista"].sum().sort_values(ascending=False)
        par = par[par > 0]
        if not par.empty:
            pdf = par.reset_index()
            pdf.columns = ["Cuenta", "Monto"]
            pdf["Acumulado"] = pdf["Monto"].cumsum() / pdf["Monto"].sum() * 100
            pdf = pdf.head(15)
            b = alt.Chart(pdf).encode(x=alt.X("Cuenta:N", sort=None, title=None))
            barras = b.mark_bar(color=MESS_BLUE).encode(
                y=alt.Y("Monto:Q", title=cur), tooltip=["Cuenta", alt.Tooltip("Monto:Q", format=",.0f")])
            linea = b.mark_line(color=ACCENT, point=True).encode(
                y=alt.Y("Acumulado:Q", scale=alt.Scale(domain=[0, 100]), title="% acumulado"))
            regla = alt.Chart(pd.DataFrame({"y": [80]})).mark_rule(color=RED, strokeDash=[5, 5]).encode(
                y=alt.Y("y:Q", scale=alt.Scale(domain=[0, 100])))
            show_chart(alt.layer(barras, linea + regla).resolve_scale(y="independent").properties(height=320))
        else:
            st.info("Sin montos en esta vista.")
    with g2:
        st.markdown("**Embudo: bruto frente a ponderado**")
        emb = dv.groupby("Fase").agg(Bruto=("Valor_Vista", "sum"), Ponderado=("Pond_Vista", "sum")).reset_index()
        if not emb.empty and emb["Bruto"].sum() > 0:
            ch_b = alt.Chart(emb).mark_bar(color="#a8c3de").encode(
                y=alt.Y("Fase:N", sort="ascending", title=None), x=alt.X("Bruto:Q", title=cur, stack=None),
                tooltip=["Fase", alt.Tooltip("Bruto:Q", format=",.0f")])
            ch_p = alt.Chart(emb).mark_bar(color=MESS_BLUE, size=10).encode(
                y=alt.Y("Fase:N", sort="ascending"), x=alt.X("Ponderado:Q", stack=None),
                tooltip=["Fase", alt.Tooltip("Ponderado:Q", format=",.0f")])
            show_chart(alt.layer(ch_b, ch_p).properties(height=320))
            st.caption("Barra clara: monto bruto. Barra oscura: monto ponderado.")
        else:
            st.info("Sin montos en esta vista.")

    st.markdown("**Forecast ponderado por área o marca**")
    fa = dv.groupby("Area")["Pond_Vista"].sum().sort_values(ascending=False).head(12).reset_index()
    fa = fa[fa["Pond_Vista"] > 0]
    if not fa.empty:
        show_chart(alt.Chart(fa).mark_bar(color=MESS_DARK).encode(
            y=alt.Y("Area:N", sort="-x", title=None), x=alt.X("Pond_Vista:Q", title="Ponderado " + cur),
            tooltip=["Area", alt.Tooltip("Pond_Vista:Q", format=",.0f")]).properties(height=300))

    # --- Ranking de cuentas clave -----------------------------------------------
    st.markdown("### Ranking de cuentas clave")
    rk = df.groupby("Cuenta").agg(
        Proyectos=("Folio", "count"), USD=("Monto_USD", "sum"), MXN=("Monto_MXN", "sum"),
        Total=("Total_USD_Eq", "sum"), Ponderado=("Ponderado", "sum"),
        Fase_Max=("Fase_Num", "max"), Dias_max=("Dias_Registro", "max"),
        En_Rojo=("Semaforo", lambda x: int((x == "Rojo").sum())),
        Sin_Cotizar=("Sin_Cotizar", "sum"),
    ).reset_index().sort_values("Total", ascending=False).reset_index(drop=True)
    total_rk = float(rk["Total"].sum())
    if total_rk > 0:
        rk["Pct"] = rk["Total"] / total_rk * 100
        rk["Pct_Acum"] = rk["Pct"].cumsum()
        rk["Clave"] = np.where((rk["Pct_Acum"] - rk["Pct"]) < 80, "Clave 80%", "")
        n80 = int((rk["Clave"] != "").sum())
        st.caption(
            str(n80) + " de " + str(len(rk)) + " cuentas concentran el 80% del valor del pipeline filtrado. "
            "Ordenado por total equivalente en USD; el ponderado usa la probabilidad por fase."
        )
    else:
        rk["Pct"], rk["Pct_Acum"], rk["Clave"] = 0.0, 0.0, ""
    rk["Fase más avanzada"] = rk["Fase_Max"].map(FASES)
    rk.insert(0, "#", range(1, len(rk) + 1))
    rk_show = rk[["#", "Cuenta", "Proyectos", "USD", "MXN", "Total", "Pct", "Pct_Acum", "Ponderado",
                  "Fase más avanzada", "Dias_max", "En_Rojo", "Sin_Cotizar", "Clave"]].rename(columns={
        "Total": "Total Eq. USD", "Pct": "% del total", "Pct_Acum": "% acum.", "Ponderado": "Ponderado USD",
        "Dias_max": "Días reg. máx.", "En_Rojo": "En rojo", "Sin_Cotizar": "Sin cotizar",
    })
    rk_sty = rk_show.style.format({
        "USD": "${:,.2f}", "MXN": "${:,.2f}", "Total Eq. USD": "${:,.0f}", "Ponderado USD": "${:,.0f}",
        "% del total": "{:.1f}%", "% acum.": "{:.1f}%",
    })
    try:
        st.dataframe(rk_sty, hide_index=True, width="stretch")
    except TypeError:
        st.dataframe(rk_sty, hide_index=True, use_container_width=True)

    # --- Cuentas con varios proyectos -------------------------------------------
    st.markdown("### Cuentas con 2 o más proyectos (candidatas a bundle)")
    cta = df.groupby("Cuenta").agg(
        Proyectos=("Folio", "count"), USD=("Monto_USD", "sum"), MXN=("Monto_MXN", "sum"),
        Total=("Total_USD_Eq", "sum"), Ponderado=("Ponderado", "sum"),
        Dias_max=("Dias_Registro", "max"),
    ).reset_index()
    cta = cta[cta["Proyectos"] >= 2].sort_values("Total", ascending=False)
    if cta.empty:
        st.info("No hay cuentas con múltiples proyectos bajo estos filtros.")
    else:
        sty = cta.rename(columns={"Total": "Total Eq. USD", "Ponderado": "Ponderado USD", "Dias_max": "Días reg. máx."}).style.format(
            {"USD": "${:,.2f}", "MXN": "${:,.2f}", "Total Eq. USD": "${:,.0f}", "Ponderado USD": "${:,.0f}"})
        try:
            st.dataframe(sty, hide_index=True, width="stretch")
        except TypeError:
            st.dataframe(sty, hide_index=True, use_container_width=True)

    # --- Detalle completo ------------------------------------------------------------
    st.markdown("### Detalle completo del pipeline")
    det = df.sort_values("Total_USD_Eq", ascending=False)
    cols_det = ["Folio", "Cliente", "Cuenta", "Descripcion", "Area", "Fase", "Monto_USD", "Monto_MXN",
                "Total_USD_Eq", "Ponderado", "Cotizacion", "Cierre_Est", "Dias_Registro", "Semaforo", "Alerta"]
    show_table(det, cols_det)
    totals_caption(det)
    exp = det[cols_det].rename(columns=DISPLAY)
    st.download_button("Descargar vista filtrada (CSV)", exp.to_csv(index=False).encode("utf-8-sig"),
                       file_name="pipeline_filtrado.csv", mime="text/csv")


# ---------------------------------------------------------------------------
# LABORATORIO IA
# ---------------------------------------------------------------------------
def _apply_quick():
    lab = st.session_state.get("lab_quick")
    qmap = st.session_state.get("lab_quick_map", {})
    if lab in qmap:
        cuenta, folio = qmap[lab]
        st.session_state["lab_cuenta"] = cuenta
        st.session_state["lab_folios_" + cuenta] = [folio]


def render_lab():
    st.markdown("### Laboratorio IA | Copiloto estratégico KAM")
    if base is None or df.empty:
        st.info("Sube el CSV y ajusta los filtros para trabajar una cuenta.")
        return
    st.caption("Trabaja con los proyectos que pasan los filtros activos del panel lateral.")

    top = df[df["Total_USD_Eq"] > 0].sort_values("Ponderado", ascending=False).head(10)
    qmap = {}
    for _, r in top.iterrows():
        qmap[r["Folio"] + " | " + r["Cuenta"] + " | " + fmt_usd(r["Total_USD_Eq"])] = (r["Cuenta"], r["Folio"])
    st.session_state["lab_quick_map"] = qmap
    st.selectbox("Atajo: proyecto prioritario (Top 10)", ["--"] + list(qmap.keys()), key="lab_quick", on_change=_apply_quick)

    cuentas = sorted(df["Cuenta"].unique())
    if st.session_state.get("lab_cuenta") not in ["--"] + cuentas:
        st.session_state.pop("lab_cuenta", None)
    cuenta = st.selectbox("Cuenta", ["--"] + cuentas, key="lab_cuenta")
    if cuenta == "--":
        return

    sub = df[df["Cuenta"] == cuenta].sort_values("Ponderado", ascending=False)
    etiquetas = {
        r["Folio"]: "Proj: " + r["Folio"] + " | Cot: " + (r["Cotizacion"] or "N/D") + " | " + str(r["Descripcion"])[:40] + " | " + r["Fase"] + " | $"
        + format(r["Monto_USD"], ",.0f") + " USD / $" + format(r["Monto_MXN"], ",.0f") + " MXN | "
        + str(r["Dias_Registro"]) + " d"
        for _, r in sub.iterrows()
    }
    folios = st.multiselect(
        "Folios a incluir (bundle)", list(etiquetas.keys()), default=list(etiquetas.keys())[:1],
        format_func=lambda x: etiquetas.get(x, x), key="lab_folios_" + cuenta,
    )
    if not folios:
        st.info("Selecciona al menos un folio.")
        return
    sel = sub[sub["Folio"].isin(folios)]

    usd, mxn = float(sel["Monto_USD"].sum()), float(sel["Monto_MXN"].sum())
    desglose = "\n".join([
        "• Cotización " + (r["Cotizacion"] or "sin número") + " (Proyecto " + r["Folio"] + "): " + str(r["Descripcion"])
        + " -> Inversión: $" + format(r["Monto_USD"], ",.2f") + " USD / $" + format(r["Monto_MXN"], ",.2f") + " MXN"
        for _, r in sel.iterrows()
    ])
    fase_n = int(sel["Fase_Num"].max())
    cierre_min = sel["Cierre_Est"].min()
    st.markdown(
        '<div class="kpi-card"><div class="kpi-label">Valor del paquete</div>'
        '<div class="kpi-value">$' + format(usd, ",.2f") + " USD | $" + format(mxn, ",.2f") + " MXN</div>"
        '<div class="kpi-sub">' + FASES[fase_n] + " | equivalente " + fmt_usd(sel["Total_USD_Eq"].sum())
        + " | ponderado " + fmt_usd(sel["Ponderado"].sum()) + "</div></div>",
        unsafe_allow_html=True,
    )

    obj_def = OBJETIVO_FASE.get(fase_n, OBJETIVOS[0])
    if bool(sel["Sin_Cotizar"].all()):
        obj_def = OBJETIVOS[0]
    c1, c2 = st.columns(2)
    with c1:
        objetivo = st.selectbox("Objetivo", OBJETIVOS, index=OBJETIVOS.index(obj_def) if obj_def in OBJETIVOS else 0,
                                key="obj_" + cuenta + "_" + str(fase_n))
        canal = st.radio("Canal", CANALES, horizontal=True)
        perfil = st.selectbox("Interlocutor", PERFILES)
    with c2:
        giro = st.text_input("Giro industrial", placeholder="Ej. Aeroespacial, automotriz")
        contexto = st.text_area("Contexto específico (opcional)", height=110,
                                placeholder="Ej. Pidió cotización pero hay resistencia por presupuesto o tiempos.")

    if st.button("Generar estrategia de cierre"):
        ctx = {
            "cuenta": cuenta, "folios": ", ".join(folios),
            "descripcion": " | ".join(sel["Descripcion"].astype(str))[:700],
            "usd": usd, "mxn": mxn, "fase": FASES[fase_n], "dias": int(sel["Dias_Registro"].max()),
            "cierre": cierre_min.strftime("%d/%m/%Y") if pd.notna(cierre_min) else "Sin fecha",
            "cotizaciones": ", ".join([c for c in sel["Cotizacion"].astype(str) if c]),
            "desglose": desglose,
            "objetivo": objetivo, "canal": canal, "perfil": perfil, "giro": giro, "contexto": contexto,
            "foco": FOCO_MEDDPICC.get(fase_n, FOCO_MEDDPICC[0]),
        }
        with st.spinner("Analizando la cuenta con MEDDPICC, SPIN y Sandler..."):
            txt, err = call_gemini(build_prompt(ctx))
        if err:
            st.error(err)
        else:
            tags = parse_tags(txt)
            st.session_state["ia_counter"] = st.session_state.get("ia_counter", 0) + 1
            st.session_state["ia_res"] = {
                "key": (cuenta, tuple(sorted(folios))), "ctx": ctx, "tags": tags, "raw": txt,
                "id": st.session_state["ia_counter"],
            }

    res = st.session_state.get("ia_res")
    if not res:
        return
    if res["key"] != (cuenta, tuple(sorted(folios))):
        st.info("Hay una estrategia generada para otra selección. Pulsa 'Generar estrategia de cierre' para esta.")
        return

    tags, ctx, rid = res["tags"], res["ctx"], str(res["id"])
    if not any(tags.values()):
        st.warning("La IA no devolvió las etiquetas esperadas. Se muestra la respuesta completa; vuelve a generar.")
        st.write(res["raw"])
        return
    if not tags["MESSAGE"]:
        st.warning("La respuesta no incluyó la etiqueta MESSAGE. Vuelve a generar.")

    st.divider()
    pa, pm, po, pl = st.tabs(["Análisis", "Mensaje", "Objeciones", "Log para SCOTT"])
    with pa:
        st.write(tags["ANALYSIS"])
    with pm:
        versiones = [clean_message(v) for v in re.split(r"\n\s*-{3,}\s*\n", tags["MESSAGE"]) if v.strip()]
        if len(versiones) >= 2:
            ver = st.radio("Versión", ["Breve", "Completa"], horizontal=True, key="ver_" + rid)
            base_msg = versiones[0] if ver == "Breve" else versiones[1]
        else:
            ver = "Unica"
            base_msg = clean_message(tags["MESSAGE"])
        mensaje = st.text_area("Mensaje (editable)", value=base_msg, height=190, key="msg_" + rid + "_" + ver)
    with po:
        st.write(tags["OBJECTIONS"])
    with pl:
        log_txt = st.text_area("Registro sugerido (editable)", value=tags["LOG"], height=110, key="log_" + rid)

    st.markdown("##### Ejecución rápida")
    a1, a2 = st.columns(2)
    with a1:
        phone = st.text_input("Teléfono WhatsApp (opcional, con lada, ej. 524421234567)", key="ph_" + rid)
        st.link_button("Abrir WhatsApp con el mensaje", whatsapp_link(mensaje, phone))
    with a2:
        d1, d2 = st.columns(2)
        with d1:
            f_cita = st.date_input("Seguimiento", value=date.today() + timedelta(days=2), key="fc_" + rid)
        with d2:
            h_cita = st.time_input("Hora", value=time(10, 0), key="hc_" + rid)
        st.link_button(
            "Crear recordatorio en Google Calendar",
            gcal_link("Seguimiento " + cuenta, ctx["objetivo"] + "\n\n" + (log_txt or ""), datetime.combine(f_cita, h_cita)),
        )
    if Document is None:
        st.warning("Instala python-docx para descargar la ficha .docx.")
    else:
        try:
            st.download_button(
                "Descargar ficha estratégica (.docx)",
                data=build_docx(ctx, tags, mensaje, log_txt),
                file_name="Estrategia_" + re.sub(r"\W+", "_", cuenta) + ".docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        except Exception as e:
            st.error("No se pudo generar el .docx: " + str(e))



# ---------------------------------------------------------------------------
# MOTOR DE MARKETING DIGITAL (apoyo a la venta)
# ---------------------------------------------------------------------------
def build_mk_prompt(ctx):
    lines = [
        "Actúa como estratega de marketing digital B2B de Grupo MESS (metrología industrial en México), apoyando a Javier Camacho, Key Account Manager.",
        "Tu trabajo es crear acciones de marketing que ayuden a cerrar ventas ya abiertas, no campañas genéricas.",
        "Lenguaje metrológico natural (trazabilidad, incertidumbre, repetibilidad, paros de línea, retrabajo), tono " + ctx["tono"].lower() + ", sin exageraciones.",
        "",
        "HECHOS VERIFICABLES DE MESS (apóyate solo en estos; no inventes certificaciones, cifras, clientes, casos de éxito ni plazos):",
    ]
    lines += ["- " + h for h in MESS_HECHOS]
    lines += [
        "",
        "BASE DE LA ACCIÓN:",
        ctx["resumen"],
        "- Objetivo de marketing: " + ctx["objetivo"],
        "- Canales solicitados: " + ", ".join(ctx["canales"]),
        "- Giro o público: " + (ctx["giro"] or "No especificado"),
        "- Evento, oferta o novedad a mencionar: " + (ctx["oferta"] or "Ninguna"),
        "",
        "REGLAS:",
        "- El contenido de LinkedIn y de la FICHA es público: nunca menciones nombres de clientes ni montos.",
        "- Si un caso de éxito no está en los hechos, no lo inventes; propón 'tema de caso a documentar'.",
        "- Aplica SPIN en los ganchos y Sandler en los cierres: da permiso de decir que no, sin presión.",
        "- Texto plano: prohibido usar tablas Markdown y asteriscos.",
        "- Para los canales no solicitados escribe únicamente: No solicitado.",
        "",
        "FORMATO: responde ÚNICAMENTE con estas siete etiquetas XML, en este orden, sin texto fuera de ellas ni bloques de código:",
        "<ESTRATEGIA> Mensaje central, audiencia, por qué ahora y cómo esta acción apoya el cierre en la fase actual (5 a 8 líneas). </ESTRATEGIA>",
        "<LINKEDIN> Una publicación lista para publicar: gancho, tres ideas, llamada a la acción y máximo cinco hashtags. </LINKEDIN>",
        "<EMAILS> Secuencia de tres correos (día 0, día 3 y día 7), cada uno con 'Asunto:' y cuerpo corto, separados por una línea con tres guiones (---). </EMAILS>",
        "<WHATSAPP> Dos mensajes cortos de seguimiento separados por una línea con tres guiones (---). </WHATSAPP>",
        "<FICHA> Texto para una ficha de una página o landing: titular, subtítulo y tres bloques (Qué resuelve, Por qué MESS, Siguiente paso). </FICHA>",
        "<CALENDARIO> Plan de 14 días: una línea por acción con el día y el responsable (KAM o Marketing). </CALENDARIO>",
        "<KPIS> Entre cuatro y seis indicadores para medir si la acción apoyó la venta (respuestas, cotizaciones reabiertas, reuniones agendadas, POs). Sin prometer resultados ni inventar cifras. </KPIS>",
    ]
    return "\n".join(lines)


def build_mk_docx(ctx, tags):
    doc = Document()
    doc.add_heading("Plan de marketing digital | Grupo MESS", 0)
    doc.add_paragraph("Fecha: " + datetime.now().strftime("%d/%m/%Y") + "\nObjetivo: " + ctx["objetivo"]
                      + "\nCanales: " + ", ".join(ctx["canales"]))
    for titulo, tag in [("Estrategia", "ESTRATEGIA"), ("Publicación de LinkedIn", "LINKEDIN"),
                        ("Secuencia de correos", "EMAILS"), ("Mensajes de WhatsApp", "WHATSAPP"),
                        ("Ficha o landing", "FICHA"), ("Calendario de 14 días", "CALENDARIO"),
                        ("Indicadores", "KPIS")]:
        doc.add_heading(titulo, level=2)
        doc.add_paragraph(tags.get(tag, "") or "Sin contenido.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def render_marketing():
    st.markdown("### Motor de Marketing Digital | Apoyo a la venta")
    if base is None or df.empty:
        st.info("Sube el CSV y ajusta los filtros para generar acciones de marketing.")
        return
    st.caption("Convierte tu pipeline en acciones de apoyo: contenido, correos, WhatsApp, ficha y calendario de 14 días.")

    modo = st.radio("Base de la acción", ["Oportunidad (cuenta y folios)", "Campaña por segmento (filtros activos)"],
                    horizontal=True, key="mk_modo")
    if modo.startswith("Oportunidad"):
        cuentas = sorted(df["Cuenta"].unique())
        if st.session_state.get("mk_cuenta") not in ["--"] + cuentas:
            st.session_state.pop("mk_cuenta", None)
        cuenta = st.selectbox("Cuenta", ["--"] + cuentas, key="mk_cuenta")
        if cuenta == "--":
            return
        sub = df[df["Cuenta"] == cuenta].sort_values("Ponderado", ascending=False)
        etiq = {r["Folio"]: "Proj: " + r["Folio"] + " | " + str(r["Descripcion"])[:50] + " | " + r["Fase"] for _, r in sub.iterrows()}
        folios = st.multiselect("Folios", list(etiq.keys()), default=list(etiq.keys())[:1],
                                format_func=lambda x: etiq.get(x, x), key="mk_folios_" + cuenta)
        if not folios:
            st.info("Selecciona al menos un folio.")
            return
        sel = sub[sub["Folio"].isin(folios)]
        fase_n = int(sel["Fase_Num"].max())
        cierre = sel["Cierre_Est"].min()
        resumen = "\n".join([
            "- Tipo: oportunidad abierta en una cuenta (el nombre solo puede usarse en correos y WhatsApp privados): " + cuenta,
            "- Equipos o servicios: " + " | ".join(sel["Descripcion"].astype(str))[:600],
            "- Fase: " + FASES[fase_n] + " | Días desde el registro: " + str(int(sel["Dias_Registro"].max())),
            "- Cierre estimado: " + (cierre.strftime("%d/%m/%Y") if pd.notna(cierre) else "Sin fecha"),
            "- Monto: " + format(sel["Monto_USD"].sum(), ",.2f") + " USD y " + format(sel["Monto_MXN"].sum(), ",.2f") + " MXN",
        ])
        firma = ("op", cuenta, tuple(sorted(folios)))
    else:
        sel = df
        fases = ", ".join(k + ": " + str(v) for k, v in sel["Fase"].value_counts().sort_index().items())
        areas = ", ".join(sel.groupby("Area")["Total_USD_Eq"].sum().sort_values(ascending=False).head(5).index)
        resumen = "\n".join([
            "- Tipo: campaña por segmento, sin nombres de cliente (" + str(len(sel)) + " proyectos de " + str(sel["Cuenta"].nunique()) + " cuentas)",
            "- Fases: " + fases,
            "- Áreas o marcas principales: " + areas,
            "- Antigüedad promedio: " + str(int(sel["Dias_Registro"].mean())) + " días | Sin cotizar: " + str(int(sel["Sin_Cotizar"].sum())),
        ])
        firma = ("seg", len(sel), round(float(sel["Total_USD_Eq"].sum())))
        st.info("Campaña sobre " + str(len(sel)) + " proyectos que pasan los filtros del panel lateral. No se envían nombres de cliente a la IA.")

    c1, c2 = st.columns(2)
    with c1:
        objetivo = st.selectbox("Objetivo de marketing", MK_OBJETIVOS)
        canales = st.multiselect("Canales", MK_CANALES, default=["LinkedIn", "Correo", "WhatsApp"])
        tono = st.selectbox("Tono", MK_TONOS)
    with c2:
        giro = st.text_input("Giro o público", placeholder="Ej. Automotriz, aeroespacial, laboratorios de calidad", key="mk_giro")
        oferta = st.text_area("Evento, oferta o novedad (opcional)", height=100, key="mk_oferta",
                              placeholder="Ej. Seguimiento posterior a un evento, capacitación, renta de equipo.")

    if st.button("Generar plan de marketing"):
        if not canales:
            st.warning("Elige al menos un canal.")
        else:
            ctx = {"resumen": resumen, "objetivo": objetivo, "canales": canales, "tono": tono, "giro": giro, "oferta": oferta}
            with st.spinner("Armando el plan de marketing..."):
                txt, err = call_gemini(build_mk_prompt(ctx))
            if err:
                st.error(err)
            else:
                st.session_state["mk_counter"] = st.session_state.get("mk_counter", 0) + 1
                st.session_state["mk_res"] = {"firma": firma, "ctx": ctx, "raw": txt, "id": st.session_state["mk_counter"],
                                              "tags": parse_tags(txt, MK_TAGS)}

    res = st.session_state.get("mk_res")
    if not res:
        return
    if res["firma"] != firma:
        st.info("Hay un plan generado para otra selección. Pulsa 'Generar plan de marketing' para esta.")
        return
    tags = {k: clean_message(v) for k, v in res["tags"].items()}
    ctx, rid = res["ctx"], str(res["id"])
    if not any(tags.values()):
        st.warning("La IA no devolvió las etiquetas esperadas. Se muestra la respuesta completa; vuelve a generar.")
        st.write(res["raw"])
        return

    st.divider()
    names = ["Estrategia", "LinkedIn", "Correos", "WhatsApp", "Ficha o landing", "Calendario 14 días", "Indicadores"]
    keys = ["ESTRATEGIA", "LINKEDIN", "EMAILS", "WHATSAPP", "FICHA", "CALENDARIO", "KPIS"]
    for tab, name, key in zip(st.tabs(names), names, keys):
        with tab:
            st.text_area(name + " (editable)", value=tags.get(key, ""), height=300, key="mk_" + key + "_" + rid)

    st.markdown("##### Ejecución rápida")
    wa_msgs = [m.strip() for m in re.split(r"\n\s*-{3,}\s*\n", tags.get("WHATSAPP", "")) if m.strip()]
    a1, a2 = st.columns(2)
    with a1:
        if wa_msgs and "No solicitado" not in wa_msgs[0]:
            pick = st.radio("Mensaje de WhatsApp", ["Mensaje 1", "Mensaje 2"][:len(wa_msgs)], horizontal=True, key="mk_wa_" + rid)
            msg = wa_msgs[0] if pick == "Mensaje 1" else wa_msgs[1]
            phone = st.text_input("Teléfono (opcional, con lada, ej. 524421234567)", key="mk_ph_" + rid)
            st.link_button("Abrir WhatsApp con el mensaje", whatsapp_link(msg, phone))
    with a2:
        d1, d2 = st.columns(2)
        with d1:
            f_acc = st.date_input("Primera acción", value=date.today() + timedelta(days=1), key="mk_fd_" + rid)
        with d2:
            h_acc = st.time_input("Hora", value=time(9, 0), key="mk_hd_" + rid)
        st.link_button("Programar en Google Calendar",
                       gcal_link("Marketing de apoyo: " + ctx["objetivo"], tags.get("CALENDARIO", ""), datetime.combine(f_acc, h_acc)))
    if Document is None:
        st.warning("Instala python-docx para descargar el plan en .docx.")
    else:
        try:
            st.download_button("Descargar plan de marketing (.docx)", data=build_mk_docx(ctx, tags),
                               file_name="Plan_marketing_MESS.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        except Exception as e:
            st.error("No se pudo generar el .docx: " + str(e))

with tab_dash:
    try:
        render_dashboard()
    except Exception as e:
        st.error("Error en el dashboard: " + str(e))
with tab_lab:
    try:
        render_lab()
    except Exception as e:
        st.error("Error en el Laboratorio IA: " + str(e))
with tab_mk:
    try:
        render_marketing()
    except Exception as e:
        st.error("Error en el motor de marketing: " + str(e))

st.markdown(
    '<div class="brand-footer">Grupo MESS | all about metrology | mess.com.mx</div>',
    unsafe_allow_html=True,
)
